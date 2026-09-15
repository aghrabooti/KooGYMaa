# -*- coding: utf-8 -*-
"""Rough PPTX -> PNG preview renderer (layout sanity check only).

Handles the subset of shapes used by build_deck.py: solid rectangles,
rounded rectangles, arrows, pictures and textboxes. Font metrics differ from
PowerPoint's, so this is for checking overlap/overflow, not pixel fidelity.
"""
import os, sys, math
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
SCALE = 110 / 914400.0  # px per EMU (110 dpi)
REG = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
_cache = {}


def font(sz, bold=False):
    key = (round(sz), bold)
    if key not in _cache:
        _cache[key] = ImageFont.truetype(BOLD if bold else REG, max(6, round(sz)))
    return _cache[key]


def px(v):
    return int(round(Emu(int(v)) * SCALE))


def shape_fill(sh):
    try:
        if sh.fill.type is not None and sh.fill.type == 1:
            c = sh.fill.fore_color
            return (c.rgb[0], c.rgb[1], c.rgb[2])
    except Exception:
        pass
    return None


def render(pptx_path, outdir):
    os.makedirs(outdir, exist_ok=True)
    prs = Presentation(pptx_path)
    W, H = px(prs.slide_width), px(prs.slide_height)
    for i, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (W, H), (12, 16, 14))
        d = ImageDraw.Draw(img, "RGBA")
        for sh in slide.shapes:
            x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
            if sh.shape_type == 13:  # picture
                try:
                    im = Image.open(sh.image.blob and __import__("io").BytesIO(sh.image.blob))
                    img.paste(im.convert("RGB").resize((max(1, w), max(1, h))), (x, y))
                except Exception:
                    d.rectangle([x, y, x + w, y + h], fill=(40, 48, 42))
                continue
            f = shape_fill(sh)
            if f is not None:
                st = str(sh.shape_type)
                if "ROUNDED" in st:
                    d.rounded_rectangle([x, y, x + w, y + h], radius=min(w, h) // 4, fill=f)
                elif "ARROW" in st:
                    d.rectangle([x, y, x + w, y + h], fill=f)
                else:
                    d.rectangle([x, y, x + w, y + h], fill=f)
            try:
                ln = sh.line
                if ln.fill.type == 1:
                    c = ln.color.rgb
                    d.rectangle([x, y, x + w, y + h], outline=(c[0], c[1], c[2]), width=1)
            except Exception:
                pass
            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            ml = px(tf.margin_left or 0); mt = px(tf.margin_top or 0)
            mr = px(tf.margin_right or 0)
            cy = y + mt
            anchor = str(tf.vertical_anchor)
            total = 0
            blocks = []
            for p in tf.paragraphs:
                text = "".join(r.text for r in p.runs)
                if not p.runs:
                    blocks.append((None, "", 0, 0, None, 1.0)); total += 8; continue
                r0 = p.runs[0]
                sz = (r0.font.size.pt if r0.font.size else 18)
                bd = bool(r0.font.bold)
                col = r0.font.color.rgb if r0.font.color and r0.font.color.type is not None else None
                col = (col[0], col[1], col[2]) if col else (240, 240, 240)
                ls = p.line_spacing or 1.0
                fnt = font(sz * SCALE * 914400 / 72 / 1.0 * 0 + sz * 110 / 72, bd)
                maxw = max(10, w - ml - mr)
                words = text.split(" ")
                lines, cur = [], ""
                for wd in words:
                    t = (cur + " " + wd).strip()
                    if d.textlength(t, font=fnt) <= maxw or not cur:
                        cur = t
                    else:
                        lines.append(cur); cur = wd
                lines.append(cur)
                lh = (fnt.size * 1.25) * (ls if ls else 1.0)
                blocks.append((fnt, lines, lh, col, p.alignment, ls))
                total += lh * len(lines) + (p.space_after.pt * 110 / 72 if p.space_after else 0)
            if "MIDDLE" in anchor:
                cy = y + (h - total) / 2
            elif "BOTTOM" in anchor:
                cy = y + h - total
            for blk in blocks:
                fnt, lines, lh, col, align, ls = blk
                if fnt is None:
                    cy += 8; continue
                for ln_ in lines:
                    tw = d.textlength(ln_, font=fnt)
                    tx = x + ml
                    if align is not None and "CENTER" in str(align):
                        tx = x + (w - tw) / 2
                    elif align is not None and "RIGHT" in str(align):
                        tx = x + w - mr - tw
                    d.text((tx, cy), ln_, font=fnt, fill=col)
                    cy += lh
        img.save(os.path.join(outdir, f"slide-{i:02d}.png"))
    print("rendered", len(prs.slides._sldIdLst), "slides to", outdir)


if __name__ == "__main__":
    render(sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "KooGYMaa.pptx"),
           sys.argv[2] if len(sys.argv) > 2 else "/tmp/deck-preview")
